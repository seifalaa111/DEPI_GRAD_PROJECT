from __future__ import annotations

import importlib.util
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Lungify_Project_Documentation.docx"
TABLE_GEOMETRY = Path(
    "C:/Users/seif alaa/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.630.12135/skills/documents/scripts/table_geometry.py"
)

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)
LIGHT_FILL = "F4F6F9"
HEADER_FILL = "F2F4F7"
BORDER = "D9DEE7"


def load_table_geometry():
    spec = importlib.util.spec_from_file_location("table_geometry", TABLE_GEOMETRY)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


table_geometry = load_table_geometry()


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_border(paragraph, *, bottom=None):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    if bottom:
        edge = OxmlElement("w:bottom")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(bottom.get("size", 10)))
        edge.set(qn("w:space"), str(bottom.get("space", 1)))
        edge.set(qn("w:color"), bottom.get("color", "000000"))
        p_bdr.append(edge)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "8")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10


def configure_running_header_footer(section):
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    header_p.paragraph_format.line_spacing = 1.0
    header_run = header_p.add_run("LUNGIFY PROJECT DOCUMENTATION")
    set_run_font(header_run, size=9, color=MUTED, bold=True)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer_p.paragraph_format.space_after = Pt(0)
    footer_run = footer_p.add_run("Research prototype only. Not for clinical diagnosis.")
    set_run_font(footer_run, size=8.5, color=MUTED, italic=True)


def add_title_block(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Lungify Project Documentation")
    set_run_font(run, size=23, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("System overview, deployment scope, technical workflow, and engineering handoff notes")
    set_run_font(run, size=13.5, color=MUTED)

    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    rows = [
        ("Document Type", "Technical project documentation"),
        ("Project", "Lungify"),
        ("Primary Scope", "CT screening and AI report generation"),
        ("Secondary Scope", "Subtype support and segmentation visualization"),
        ("Date", "July 3, 2026"),
        ("Version", "1.1"),
    ]
    for idx, (label, value) in enumerate(rows):
        left, right = table.rows[idx].cells
        for cell in (left, right):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
        shade_cell(left, HEADER_FILL)
        lp = left.paragraphs[0]
        lp.paragraph_format.space_after = Pt(2)
        lrun = lp.add_run(label)
        set_run_font(lrun, size=10.5, color=INK, bold=True)
        rp = right.paragraphs[0]
        rp.paragraph_format.space_after = Pt(2)
        rrun = rp.add_run(value)
        set_run_font(rrun, size=10.5, color=INK)

    table_geometry.apply_table_geometry(table, [2000, 7360], table_width_dxa=9360, indent_dxa=120)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(12)
    rule.paragraph_format.space_after = Pt(10)
    set_paragraph_border(rule, bottom={"color": "AEB9C8", "size": 10, "space": 1})


def add_callout(doc: Document, title: str, body: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    shade_cell(cell, LIGHT_FILL)
    set_cell_border(cell, color="C7D1DF")

    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(4)
    r1 = p1.add_run(title)
    set_run_font(r1, size=11, color=INK, bold=True)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)

    table_geometry.apply_table_geometry(table, [9360], table_width_dxa=9360, indent_dxa=120)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_number(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_path_paragraph(doc: Document, label: str, path: str, purpose: str):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=11, color=INK, bold=True)
    r2 = p.add_run(path)
    set_run_font(r2, name="Courier New", size=10, color=DARK_BLUE)
    r3 = p.add_run(f" - {purpose}")
    set_run_font(r3, size=11, color=INK)


def add_code_block(doc: Document, lines: list[str]):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, "FAFBFC")
    set_cell_border(cell, color="D6DDE6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, name="Courier New", size=9.5, color=INK)
        if idx < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)
    table_geometry.apply_table_geometry(table, [9360], table_width_dxa=9360, indent_dxa=120)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        cell = header_cells[idx]
        shade_cell(cell, HEADER_FILL)
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(header)
        set_run_font(run, size=10, color=INK, bold=True)

    for row_idx, row_data in enumerate(rows, start=1):
        for col_idx, value in enumerate(row_data):
            cell = table.rows[row_idx].cells[col_idx]
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(value)
            if "/" in value or value.startswith("lungify-") or value.startswith("GET ") or value.startswith("POST "):
                set_run_font(run, name="Courier New", size=9.5, color=INK)
            else:
                set_run_font(run, size=10, color=INK)

    table_geometry.apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120)


def build_document():
    doc = Document()
    style_document(doc)
    configure_running_header_footer(doc.sections[0])

    doc.core_properties.title = "Lungify Project Documentation"
    doc.core_properties.subject = "Technical project documentation"
    doc.core_properties.author = "OpenAI Codex"

    add_title_block(doc)
    add_callout(
        doc,
        "Segmentation Positioning",
        "Segmentation should be presented as a supplementary visualization capability. "
        "The primary product deliverable is CT screening plus AI report generation. "
        "Segmentation supports interpretation when available, but it is not the core output, "
        "main medical claim, or primary engineering success criterion of the system.",
    )

    doc.add_heading("1. Executive Summary", level=1)
    add_body(
        doc,
        "Lungify is a research-oriented medical AI web project focused on lung CT screening. "
        "It combines a Next.js frontend, a FastAPI backend, and exported model artifacts into a single experience "
        "that allows a user to upload a de-identified CT study and receive an AI-generated report.",
    )
    add_body(
        doc,
        "From an engineering perspective, the project is best understood as a web application with a real model-backed inference service, "
        "a fallback preview layer, and a user-facing report workspace. The most important outcome of the system is the final assessment plus the readable report body.",
    )
    add_body(
        doc,
        "The platform includes segmentation output, but that capability should be framed as a supplementary visual aid. "
        "The core project scope is CT screening and AI report generation, not segmentation-first analysis.",
    )

    doc.add_heading("2. Project Goals and Problem Statement", level=1)
    add_body(
        doc,
        "The project addresses a practical demonstration problem: how to package a medical AI pipeline into a usable website that engineers and evaluators can inspect, run, and assess. "
        "Rather than exposing raw model notebooks alone, the project wraps the model outputs inside a deployable product structure.",
    )
    add_body(doc, "The main goals of Lungify are:")
    for item in (
        "Provide a web-accessible interface for CT screening demonstration.",
        "Connect a real backend to exported model artifacts rather than relying only on screenshots or offline notebooks.",
        "Return a structured result that is understandable to a reviewer, not just a raw tensor or class index.",
        "Stay usable in demo environments even when the live backend is not available.",
        "Show a full-stack engineering integration path from upload through inference to report rendering.",
    ):
        add_bullet(doc, item)

    doc.add_heading("3. Core Product Scope", level=1)
    add_body(doc, "The main workflow of the product is:")
    for item in (
        "Accept a de-identified CT scan as a ZIP containing a DICOM series.",
        "Run the binary screening model through the backend inference service.",
        "Optionally enrich the result with advanced follow-up signals when applicable.",
        "Generate a structured AI report with summary, recommendation, and disclaimer.",
        "Return the result to the frontend for display in a clean report workspace.",
    ):
        add_number(doc, item)
    add_body(
        doc,
        "This workflow is the central engineering deliverable and should be the main framing used in technical reviews, project submissions, and demonstrations.",
    )

    doc.add_heading("4. Segmentation Scope and Positioning", level=1)
    add_body(
        doc,
        "Segmentation is part of the project, but it should not be treated as the product's central claim. "
        "The project remains valid and functional even when segmentation is unavailable, because the core product objective is report-driven screening.",
    )
    add_body(doc, "The correct positioning for segmentation is:")
    for item in (
        "A supplementary visualization and localization aid.",
        "A prototype enhancement layered on top of the primary screening workflow.",
        "A supporting result that helps interpretation when available.",
        "A non-primary feature that must not replace the final assessment and report as the main output.",
    ):
        add_bullet(doc, item)
    add_body(doc, "In the current implementation, segmentation is limited by the following product rules:")
    for item in (
        "It is tied to the advanced follow-up model pipeline.",
        "It is only attempted when the binary screening result is malignant.",
        "It depends on the advanced model bundle loading successfully at runtime.",
        "It is described by the backend itself as prototype localization rather than production-grade lesion mapping.",
    ):
        add_bullet(doc, item)

    doc.add_heading("5. System Architecture", level=1)
    add_table(
        doc,
        ["Component", "Technology / Target", "Primary Responsibility"],
        [
            ["Frontend", "Next.js / Vercel", "Public website pages, demo workspace, internal API proxy routes, and result rendering."],
            ["Backend", "FastAPI / Python", "Validation, ZIP extraction, model loading, inference execution, and response construction."],
            ["Model Artifacts", "Git LFS weights", "Pipeline A binary screening model and Pipeline B advanced follow-up bundle."],
            ["Preview Layer", "Frontend fallback mode", "Keeps the product usable when a real backend is not configured or reachable."],
        ],
        [1500, 2400, 5460],
    )
    add_body(
        doc,
        "The architecture is intentionally split into deployable units. This separation allows the frontend to remain lightweight and user-friendly, "
        "while the backend holds the heavier inference logic, validation rules, and model lifecycle concerns.",
    )

    doc.add_heading("6. Frontend Application", level=1)
    add_body(
        doc,
        "The frontend is responsible for the product experience rather than the model logic itself. "
        "It provides the landing pages, the demo upload flow, and the report viewer that presents the backend response.",
    )
    add_body(doc, "Main frontend responsibilities include:")
    for item in (
        "Hosting public-facing website pages such as the homepage, model page, demo page, and team page.",
        "Accepting file uploads in the demo interface.",
        "Routing user requests through internal API endpoints.",
        "Displaying backend health and model mode information to the user.",
        "Rendering the final report, warnings, and segmentation images when present.",
    ):
        add_bullet(doc, item)
    add_path_paragraph(doc, "Frontend root", "lungify-frontend/", "Website pages, demo client, and API proxy routes.")
    add_path_paragraph(doc, "Frontend demo UI", "lungify-frontend/components/DemoClient.tsx", "Upload interface and result workspace.")
    add_path_paragraph(doc, "Frontend predict route", "lungify-frontend/app/api/predict/route.ts", "Routes uploads toward the backend or preview mode.")
    add_path_paragraph(doc, "Frontend health route", "lungify-frontend/app/api/health/route.ts", "Used for backend status and mode reporting.")
    add_path_paragraph(doc, "Frontend model-info route", "lungify-frontend/app/api/model-info/route.ts", "Used to surface model details and limitations.")

    doc.add_heading("7. Backend Application", level=1)
    add_body(
        doc,
        "The backend is the operational core of the system. It accepts uploads, validates them, prepares DICOM data, loads model artifacts, "
        "runs inference, and shapes the final response into a frontend-friendly schema.",
    )
    add_body(doc, "Main backend responsibilities include:")
    for item in (
        "Upload classification and validation.",
        "Safe ZIP extraction for DICOM series handling.",
        "Model registry initialization and health reporting.",
        "Execution of binary screening and advanced follow-up inference.",
        "AI report assembly and warning propagation.",
    ):
        add_bullet(doc, item)
    add_path_paragraph(doc, "Backend root", "lungify-backend/", "Inference service, validation, tests, and model integration.")
    add_path_paragraph(doc, "Backend entry", "lungify-backend/app/main.py", "Starts the FastAPI application.")
    add_path_paragraph(doc, "Prediction routes", "lungify-backend/app/api/routes.py", "Implements health, model info, and predict endpoints.")
    add_path_paragraph(doc, "Model registry", "lungify-backend/app/services/model_registry.py", "Loads and exposes model availability.")
    add_path_paragraph(doc, "Validation utilities", "lungify-backend/app/utils/validation.py", "Classifies uploads and protects extraction logic.")
    add_path_paragraph(doc, "Report builder", "lungify-backend/app/services/report_builder.py", "Constructs the final response body.")

    doc.add_heading("8. Model Assets and Inference Pipelines", level=1)
    add_body(
        doc,
        "The project contains real exported model artifacts tracked with Git LFS. These artifacts are separated into a primary binary screening path and an advanced follow-up path.",
    )
    add_table(
        doc,
        ["Pipeline", "Purpose", "Runtime Role"],
        [
            ["Pipeline A", "Binary screening", "Primary classifier used to determine the main final assessment."],
            ["Pipeline B", "Subtype support and segmentation", "Secondary follow-up path used only when the binary result indicates malignancy and the advanced model is available."],
        ],
        [1800, 3000, 4560],
    )
    add_path_paragraph(doc, "Binary artifact", "lungify-backend/weights/pipeline_a_model.pkl", "Primary screening artifact.")
    add_path_paragraph(doc, "Advanced artifacts", "lungify-backend/weights/pipeline_b_models/", "Folded advanced follow-up model bundle.")
    add_body(
        doc,
        "This split is important for documentation and evaluation because it reinforces that the binary screening path is the default and primary result source. "
        "The advanced path should be treated as conditional enrichment, not the base contract of the product.",
    )

    doc.add_heading("9. Functional Request Flow", level=1)
    doc.add_heading("Real Inference Flow", level=2)
    for item in (
        "The user uploads a ZIP file containing a de-identified DICOM series.",
        "The frontend sends the request to its own /api/predict route.",
        "If LUNGIFY_BACKEND_URL is configured, the frontend proxies the request to the FastAPI backend.",
        "The backend classifies the upload and validates whether it is suitable for real inference.",
        "The backend safely extracts the ZIP contents into a temporary working directory.",
        "Pipeline A executes the main binary screening step.",
        "If the case is malignant and advanced models are loaded, Pipeline B may add subtype support and segmentation output.",
        "The backend assembles a structured response including final assessment, report, warnings, and optional images.",
        "The frontend renders the result in the report workspace.",
    ):
        add_number(doc, item)

    doc.add_heading("Preview and Sample Flow", level=2)
    add_body(
        doc,
        "The product also supports a non-real fallback mode so the website remains usable for demonstrations and product review even when the live backend is not connected.",
    )
    for item in (
        "If the backend is absent, the frontend can return a built-in preview or sample response.",
        "Single .dcm, .png, .jpg, and .jpeg uploads are preview-only in the current implementation.",
        "Preview mode is useful for interface walkthroughs, but it is not equivalent to real DICOM ZIP inference.",
        "If a backend is configured but unreachable, the frontend returns an explicit error instead of silently pretending that real inference succeeded.",
    ):
        add_bullet(doc, item)

    doc.add_heading("10. Input Contract", level=1)
    add_table(
        doc,
        ["Input Type", "Mode", "Meaning"],
        [
            [".zip with DICOM series", "Real", "Expected input for actual backend inference."],
            [".dcm", "Preview", "Single-file preview only; not full study inference."],
            [".png / .jpg / .jpeg", "Preview", "Interface preview mode only; used for demo behavior."],
        ],
        [2000, 1400, 5960],
    )
    add_body(
        doc,
        "The distinction between real and preview inputs is important in evaluation. "
        "The production-like path should always be described as ZIP-based DICOM inference.",
    )

    doc.add_heading("11. Output Contract", level=1)
    add_body(
        doc,
        "The backend returns a structured response designed for the frontend rather than a raw model dump. "
        "This improves usability and makes the output easier to review in demonstrations.",
    )
    add_table(
        doc,
        ["Response Section", "Purpose", "Priority"],
        [
            ["final_assessment", "Primary classification result including label, risk level, confidence, and source.", "Highest"],
            ["subtype_suggestion", "Supporting advanced follow-up classification information.", "Secondary"],
            ["segmentation", "Supplementary image-based visualization data when available.", "Secondary"],
            ["report", "Human-readable summary, recommendation, and disclaimer.", "Highest"],
            ["warnings", "Runtime and model-availability notes.", "Support"],
        ],
        [2200, 5200, 1960],
    )
    add_body(
        doc,
        "For presentations and engineering handoff, the most important parts of the output are final_assessment and report. "
        "Those two areas best represent the product's main purpose.",
    )

    doc.add_heading("12. API Endpoints", level=1)
    add_table(
        doc,
        ["Endpoint", "Method", "Purpose"],
        [
            ["GET /", "GET", "Basic service status endpoint."],
            ["GET /health", "GET", "Checks backend and model readiness."],
            ["GET /model-info", "GET", "Exposes architecture, metrics, classes, and limitations."],
            ["POST /predict", "POST", "Handles real upload inference and preview fallback logic."],
            ["POST /predict/sample", "POST", "Returns a sample response for demo mode."],
        ],
        [2700, 1200, 5460],
    )
    add_body(
        doc,
        "The frontend talks to its own internal API routes first. Those routes either forward the request to the real backend or respond locally in preview mode depending on deployment configuration.",
    )

    doc.add_heading("13. Validation, Safety, and Runtime Controls", level=1)
    add_body(
        doc,
        "Because the product accepts uploaded medical imaging data, the backend includes explicit validation and runtime guards around file handling and model execution.",
    )
    for item in (
        "Upload classification distinguishes between real ZIP inference and preview-only file types.",
        "Safe ZIP extraction reduces risk from malformed or unsafe archive contents.",
        "Model health endpoints make backend readiness visible before a user runs inference.",
        "Warnings are returned to the frontend when advanced models are unavailable.",
        "The project includes explicit disclaimers that it is a research prototype and not a clinical diagnostic tool.",
    ):
        add_bullet(doc, item)

    doc.add_heading("14. Repository Structure and Important Files", level=1)
    add_path_paragraph(doc, "Root README", "README.md", "Top-level project summary.")
    add_path_paragraph(doc, "Markdown documentation", "PROJECT_DOCUMENTATION.md", "Text version of the project documentation.")
    add_path_paragraph(doc, "Frontend homepage", "lungify-frontend/app/page.tsx", "Main public-facing landing page.")
    add_path_paragraph(doc, "Frontend model page", "lungify-frontend/app/model/page.tsx", "Model explanation and results page.")
    add_path_paragraph(doc, "Backend schemas", "lungify-backend/app/schemas.py", "Response schema definitions.")
    add_path_paragraph(doc, "Backend service logic", "lungify-backend/app/services/", "Inference, preprocessing, DICOM loading, and report building.")
    add_path_paragraph(doc, "Backend tests", "lungify-backend/tests/", "API, preprocessing, and validation test coverage.")

    doc.add_heading("15. Local Development", level=1)
    add_body(doc, "Recommended backend startup:")
    add_code_block(
        doc,
        [
            "cd lungify-backend",
            "pip install -r requirements.txt",
            "uvicorn app.main:app --host 0.0.0.0 --port 7860",
        ],
    )
    add_body(doc, "Recommended frontend startup:")
    add_code_block(
        doc,
        [
            "cd lungify-frontend",
            "npm install",
            "npm run dev",
        ],
    )
    add_body(doc, "Recommended local environment variable:")
    add_code_block(
        doc,
        [
            "LUNGIFY_BACKEND_URL=http://localhost:7860",
        ],
    )

    doc.add_heading("16. Deployment Model", level=1)
    add_body(
        doc,
        "The deployment design is intentionally split so the frontend can be hosted separately from the inference backend. "
        "This makes the user-facing site easier to serve while keeping model execution isolated.",
    )
    add_table(
        doc,
        ["Layer", "Deployment Target", "Notes"],
        [
            ["Frontend", "Vercel", "Uses internal /api routes and can stay usable even without a connected real backend."],
            ["Backend", "FastAPI host or Docker-capable service", "Runs model loading, validation, and inference logic."],
            ["Weights", "Git LFS backed repository artifacts", "Must be available to the runtime environment for real inference."],
        ],
        [1800, 2500, 5060],
    )
    add_body(doc, "Deployment behavior can be summarized as follows:")
    for item in (
        "With backend configured and reachable: real DICOM ZIP inference is available.",
        "Without backend configured: the site stays operational in preview and sample mode.",
        "With backend configured but unreachable: requests for real inference fail explicitly rather than silently succeeding.",
    ):
        add_bullet(doc, item)

    doc.add_heading("17. Environment Variables and Configuration", level=1)
    add_table(
        doc,
        ["Variable", "Layer", "Purpose"],
        [
            ["LUNGIFY_BACKEND_URL", "Frontend", "Base URL used by the frontend's internal API routes to reach the real backend."],
            ["LUNGIFY_CORS_ORIGINS", "Backend", "Controls which frontend origins are allowed to access the backend."],
            ["LUNGIFY_DEVICE", "Backend", "Selects runtime device behavior such as CPU."],
            ["LUNGIFY_ENV", "Backend", "Identifies the current runtime environment."],
        ],
        [2600, 1200, 5560],
    )
    add_body(
        doc,
        "A production-style deployment should explicitly set these variables instead of relying on implicit defaults.",
    )

    doc.add_heading("18. Testing and Verification", level=1)
    add_body(
        doc,
        "The repository includes both backend and frontend verification paths. "
        "At minimum, evaluators should confirm that backend tests pass, the frontend builds successfully, and the real backend path is distinguishable from preview mode.",
    )
    add_code_block(
        doc,
        [
            "cd lungify-backend",
            "python -m pytest -q",
            "",
            "cd ../lungify-frontend",
            "npm run build",
        ],
    )
    add_body(doc, "Suggested verification checklist:")
    for item in (
        "Confirm /health returns backend status and model mode information.",
        "Confirm /model-info exposes documented limitations and classes.",
        "Confirm ZIP uploads follow the real inference path when the backend is live.",
        "Confirm preview-only uploads are clearly identified as preview mode.",
        "Confirm the frontend UI displays warnings rather than hiding backend issues.",
    ):
        add_bullet(doc, item)

    doc.add_heading("19. Known Limitations", level=1)
    add_body(
        doc,
        "Like any prototype system, Lungify includes functional and product limitations that should be stated clearly during review.",
    )
    for item in (
        "The project is a research prototype and not a certified medical device.",
        "It is not intended for clinical diagnosis or replacement of radiology review.",
        "Real inference expects a de-identified DICOM ZIP input, not arbitrary single-image uploads.",
        "Advanced subtype support and segmentation depend on the advanced model bundle being loaded successfully.",
        "Segmentation is a supplementary output and should not be interpreted as the main project result.",
        "Preview mode is useful for demonstrations but does not represent the full real inference pipeline.",
    ):
        add_bullet(doc, item)

    doc.add_heading("20. Evaluation Framing and Handoff Notes", level=1)
    add_body(
        doc,
        "When this project is handed to engineers or evaluators, the strongest framing is to describe it as a full-stack CT screening and AI report platform with a real backend integration path.",
    )
    add_body(doc, "Recommended talking points:")
    for item in (
        "The project includes a real frontend, a real backend, and real model artifacts.",
        "The core deliverable is screening plus structured report generation.",
        "The system distinguishes between real inference mode and preview mode rather than hiding the difference.",
        "Segmentation exists as a supplementary prototype capability for visual support.",
        "The repository structure is suitable for engineering review because the frontend, backend, models, and tests are clearly separated.",
    ):
        add_bullet(doc, item)
    add_body(
        doc,
        "Recommended summary sentence: "
        "\"Lungify is a CT screening and AI report generation platform with a production-style frontend and backend architecture; segmentation is a supplementary visual support capability layered on top of the core workflow.\"",
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    build_document()
    print(OUTPUT)
